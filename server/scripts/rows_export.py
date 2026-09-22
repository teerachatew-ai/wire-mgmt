# -*- coding: utf-8 -*-
# Export ตารางบนหน้าจอ (array of objects) เป็นไฟล์ .xlsx — ใช้ร่วมกันทุกปุ่ม "Export Excel" ในระบบ
# เดิมสร้างไฟล์ฝั่งเบราว์เซอร์แล้วดาวน์โหลดผ่าน blob ซึ่ง Safari บน macOS เปิดไม่ได้
# ("Safari ไม่สามารถเปิดหน้าเว็บได้") จึงย้ายมาสร้างที่ server แล้วส่งเป็นไฟล์แนบตรงๆ
# Usage: python rows_export.py <data.json> <out.xlsx>
import sys, json, warnings
warnings.simplefilter("ignore")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

dataf, out = sys.argv[1], sys.argv[2]
d = json.load(open(dataf, encoding="utf-8-sig"))
rows = d.get("rows") or []
sheet = (d.get("sheet") or "Sheet1")[:31] or "Sheet1"

FONT = "TH SarabunPSK"
wb = Workbook()
ws = wb.active
ws.title = sheet

headers = []
for r in rows:
    for k in r.keys():
        if k not in headers:
            headers.append(k)

NAVY = "1E3A5F"
thin = Side(style="thin", color="D8DEE9")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

for ci, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=ci, value=h)
    c.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box
ws.row_dimensions[1].height = 24

for ri, r in enumerate(rows, start=2):
    for ci, h in enumerate(headers, start=1):
        v = r.get(h)
        if isinstance(v, bool):
            v = "ใช่" if v else "ไม่"
        elif isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False)
        c = ws.cell(row=ri, column=ci, value=v)
        c.font = Font(name=FONT, size=13)
        c.border = box
        if isinstance(v, (int, float)):
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '#,##0.##'

# ความกว้างคอลัมน์คร่าวๆ ตามความยาวข้อความ (กันตัวหนังสือถูกตัด) — เพดานเท่าของเดิมฝั่งเบราว์เซอร์
for ci, h in enumerate(headers, start=1):
    longest = max([len(str(h))] + [len(str(r.get(h) or "")) for r in rows])
    ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(max(longest + 2, 10), 45)

ws.freeze_panes = "A2"
if headers and rows:
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{len(rows) + 1}"

wb.save(out)
